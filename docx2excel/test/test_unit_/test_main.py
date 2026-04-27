import os
import tempfile
from unittest.mock import patch
from main import main


def create_test_docx(path):
    from docx import Document
    doc = Document()
    table = doc.add_table(rows=2, cols=2)

    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Age"
    table.cell(1, 0).text = "Alice"
    table.cell(1, 1).text = "18"

    doc.save(path)


@patch("main.select_word_files")
def test_main_flow(mock_select):
    with tempfile.TemporaryDirectory() as tmp_dir:
        word_path = os.path.join(tmp_dir, "test.docx")
        create_test_docx(word_path)

        mock_select.return_value = [word_path]

        # ✅ 关键：清空 sys.argv
        with patch("sys.argv", ["main.py"]):

            with patch("main.EXCEL_OUTPUT_DIR", tmp_dir):
                main()

        # 检查是否生成 Excel
        files = os.listdir(tmp_dir)

        print("生成文件：", files)

        assert any(f.endswith(".xlsx") for f in files)